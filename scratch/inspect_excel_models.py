"""
inspect_excel_models.py
"""
import openpyxl, re

excel_path = r'C:\Users\User\Documents\Coding\TechFlow Data Organizer\nanoreview_master.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['Laptops']
headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]

for r in range(2, sheet.max_row + 1):
    model = str(sheet.cell(row=r, column=headers.index('Model')+1).value or '')
    year = str(sheet.cell(row=r, column=headers.index('Date Launched')+1).value or '')
    cpu = str(sheet.cell(row=r, column=headers.index('CPU name')+1).value or '')
    print(f"Row {r:2d}: Model={model:<35} | Year={year:<6} | CPU={cpu[:35]}")
