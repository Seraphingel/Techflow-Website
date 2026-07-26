"""
test_model_cleanups.py
"""

import openpyxl, re

excel_path = r'C:\Users\User\Documents\Coding\TechFlow Data Organizer\nanoreview_master.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['Laptops']
headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]

rows = []
for r in range(2, sheet.max_row + 1):
    row_dict = {}
    for c in range(1, sheet.max_column + 1):
        val = sheet.cell(row=r, column=c).value
        if val in ['-', 'None', None]:
            val = ''
        else:
            val = str(val).strip()
        row_dict[headers[c-1]] = val
    rows.append(row_dict)

def clean_model_title(raw_model):
    if not raw_model: return ""
    # Strip any (...) from raw_model
    cleaned = re.sub(r'\s*\([^)]*\)', '', raw_model).strip()
    return cleaned

def clean_cpu_string(cpu):
    if not cpu: return ""
    # Clean up iUltra -> Ultra in raw string if present
    cpu = cpu.replace('iUltra', 'Ultra')
    return cpu.strip()

def shorten_cpu(cpu):
    if not cpu: return ""
    cpu = clean_cpu_string(cpu)
    # Remove Qualcomm, keep Snapdragon
    cpu = cpu.replace('Qualcomm Snapdragon', 'Snapdragon').replace('Qualcomm ', 'Snapdragon ')
    
    # Handle iUltra / Ultra
    if 'Intel Core Ultra' in cpu:
        cpu = 'Ultra ' + cpu.split('Intel Core Ultra')[-1].strip()
    elif 'Intel Core i' in cpu:
        cpu = 'i' + cpu.split('Intel Core i')[-1].strip()
    elif 'Intel Core ' in cpu:
        cpu = 'i' + cpu.split('Intel Core ')[-1].strip()
    elif 'iUltra' in cpu:
        cpu = 'Ultra ' + cpu.split('iUltra')[-1].strip()
    elif 'Ultra' in cpu and not cpu.startswith('Ultra'):
        cpu = 'Ultra ' + cpu.split('Ultra')[-1].strip()
        
    cpu = cpu.replace('AMD Ryzen ', 'R')
    cpu = re.sub(r'\s+Mobile.*', '', cpu)
    return cpu.strip()

processed_keys = set()
valid_models = []
skipped_no_date = []
duplicates_skipped = []

for idx, row in enumerate(rows, 1):
    brand = row.get('Brand', 'ASUS').strip()
    raw_model = row.get('Model', '').strip()
    year = row.get('Date Launched', '').strip()

    if not year or not raw_model:
        skipped_no_date.append((raw_model, year))
        continue

    clean_base_title = clean_model_title(raw_model)
    dedup_key = (brand.upper(), clean_base_title.lower(), str(year))

    if dedup_key in processed_keys:
        duplicates_skipped.append((clean_base_title, year))
        continue
    processed_keys.add(dedup_key)

    cpus = [c.strip() for c in row.get('CPU name', '').split('\n') if c.strip()]
    shortened_cpus = [shorten_cpu(c) for c in cpus]

    valid_models.append((clean_base_title, year, shortened_cpus[:2]))

print(f"Total valid models: {len(valid_models)}")
print(f"Skipped no date: {len(skipped_no_date)}")
print(f"Duplicates skipped: {len(duplicates_skipped)}")
print("\nSample cleaned model names & shortened CPUs:")
for m in valid_models[:30]:
    print(f"  - {m[0]} ({m[1]}) | CPUs: {m[2]}")
